"""
extractors/lar.py — La Rioja.

PDF `funcional_economico.pdf` (Detalle Gastos Funcional / Económico).  Es una
tabla compacta con 4 niveles jerárquicos en columnas — Grupo, Función,
Subfunción, Programa — y 9 columnas de capítulo + Total al final.  El texto
del PDF tiene varios artefactos (la cifra cero aparece como "o" minúscula, y
algunos rótulos aparecen separados por espacios — "AL TA DIRECCIÓN").

Estrategia:
  1. Si hay un XLSX/CSV adjunto en `fuentes/raw/lar/<año>/partidas*.xlsx|csv`
     con columnas (codigo, denominacion, importe_eur), gana sobre el PDF.
  2. Si no, intenta primero tablas de resumen con códigos `NNNN - programa`.
  3. Como fallback, parsea el texto del PDF con **pdfplumber** y reconstruye
     los códigos jerárquicos G.F.SF.P (e.g. `3.1.2.1` = Sanidad → Hospitales →
     Atención primaria) usando una máquina de estado sobre las líneas.

Códigos lar: 4 niveles `<Grupo>.<Función>.<Subfunción>.<Programa>` separados
por punto (formato compatible con PROG_CODIGO_RE en `_common.base`).
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from .._common import base
from .._common.base import make_row, ExtractionResult, parse_eur


_RE_LEVEL = re.compile(r"^\s*(\d+)\s*-\s*(.+?)\s*$")
_RE_CODE4 = re.compile(r"^\s*(?P<code>\d{4})\s*-\s*(?P<rest>.*)$")
_RE_ORG_PROGRAMA = re.compile(
    r"^\s*Programa\s+(?P<code>\d{4})\s+"
    r"(?P<denom>.+?)\s+"
    r"(?P<importe>\d{1,3}(?:\.\d{3})+|\d+)\s*$"
)
_RE_AMOUNT = re.compile(r"(?<!\w)-?(?:\d{1,3}(?:\.\d{3})+|\d+|[oO])(?!\w)")
_RE_SEGMENT = re.compile(r"(?<!\d)(\d)\s*(?:[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]\s*){0,8}-\s*")
_RE_INFORME_SEGMENT = re.compile(r"(?<!\d)(\d)\s+(?=[A-ZÁÉÍÓÚÜÑ])")


def _iter_input_pdfs(input_path: Path, anio: int | None = None) -> list[Path]:
    if input_path.is_dir():
        if anio == 2023:
            detalle = input_path / "detalle_gastos_funcional_economico.pdf"
            if detalle.exists():
                return [detalle]
        if anio == 2015:
            main_pdfs = sorted(p for p in input_path.glob("*.pdf")
                               if p.name != "funcional_economico_fragmento.pdf")
            if main_pdfs:
                return main_pdfs
        fragmentos = sorted((input_path / "fragmentos").glob("*.pdf"))
        if fragmentos:
            return fragmentos
        return sorted(input_path.glob("*.pdf"))
    return [input_path]


def _from_xlsx(input_path: Path, anio: int) -> list[dict]:
    candidatos = list(input_path.parent.glob("partidas*.xlsx")) + \
                 list(input_path.parent.glob("partidas*.csv"))
    if not candidatos:
        return []
    src = candidatos[0]
    rows: list[dict] = []
    if src.suffix.lower() == ".csv":
        with src.open("r", encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                if not r.get("codigo"):
                    continue
                rows.append(make_row(pagina=None, codigo=r["codigo"],
                                     denominacion=r.get("denominacion", ""),
                                     importe=parse_eur(r.get("importe_eur", "0")),
                                     anio=anio))
    else:
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError:
            return []
        wb = load_workbook(filename=str(src), data_only=True, read_only=True)
        ws = wb.active
        cells = list(ws.iter_rows(values_only=True))
        if not cells:
            return []
        header = [str(c).strip().lower() if c else "" for c in cells[0]]
        try:
            ic = header.index("codigo")
            id_ = header.index("denominacion")
            ii = header.index("importe_eur")
        except ValueError:
            return []
        for r in cells[1:]:
            if r[ic] is None:
                continue
            rows.append(make_row(pagina=None, codigo=str(r[ic]),
                                 denominacion=str(r[id_] or ""),
                                 importe=parse_eur(r[ii]) if r[ii] else float("nan"),
                                 anio=anio))
    return rows


def _norm_label(s: str) -> str:
    """Colapsa espacios y une trozos OCR como 'AL TA DIRECCIÓN' → 'ALTA DIRECCIÓN'."""
    s = (s or "").strip().replace("\n", " ")
    s = re.sub(r"\s+", " ", s)
    # Une 'AL TA' → 'ALTA', 'RE GIMEN' → 'RÉGIMEN' (heurística mínima)
    s = re.sub(r"\bAL TA\b", "ALTA", s)
    s = re.sub(r"\bRE GIMEN\b", "RÉGIMEN", s)
    return s


def _to_int_eur(s: str) -> int | None:
    s = (s or "").strip().replace(".", "").replace(" ", "")
    # El texto OCR de "0" aparece como "o" o "O"
    s = s.replace("o", "0").replace("O", "0")
    sign = -1 if s.startswith("-") else 1
    s = s.lstrip("-")
    return sign * int(s) if s.isdigit() and len(s) >= 3 else None


def _is_total_label(s: str) -> bool:
    return (s or "").strip().lower().startswith("total")


def _amounts(s: str) -> list[int]:
    vals: list[int] = []
    for m in _RE_AMOUNT.finditer(s):
        val = _to_int_eur(m.group(0))
        if val is not None:
            vals.append(val)
    return vals


def _amount_matches(s: str):
    vals = []
    for m in _RE_AMOUNT.finditer(s):
        val = _to_int_eur(m.group(0))
        if val is not None:
            vals.append((m, val))
    return vals


def _row_total(vals: list[int]) -> int | None:
    if len(vals) >= 10:
        return vals[-1]
    if len(vals) == 9:
        return sum(vals)
    if len(vals) >= 2:
        return vals[-1]
    return None


def _informe_segments(head: str) -> list[tuple[str, str]]:
    """Segmenta cabeceras `1 SANIDAD 2 HOSPITALES 1 ATENCION ...`."""
    starts = list(_RE_INFORME_SEGMENT.finditer(head))
    segments: list[tuple[str, str]] = []
    for i, m in enumerate(starts):
        end = starts[i + 1].start() if i + 1 < len(starts) else len(head)
        segments.append((m.group(1), _norm_label(head[m.end():end])))
    return segments


def _code4_to_dotted(code: str) -> str:
    return ".".join(code)


def _reverse_pdf_token(s: str) -> str:
    s = s[::-1].strip()
    # Artefactos de codificación vistos en el BOLR 2015.
    s = s.replace(")112:dic(", "Ñ")
    s = s.replace(")902:dic(", "Ñ")
    s = s.replace(")342:dic(", "Ó")
    s = s.replace(")732:dic(", "í")
    s = s.replace(")812:dic(", "Á")
    s = s.replace(")102:dic(", "É")
    return _norm_label(s)


def _page_text(page) -> str:
    try:
        return page.extract_text(x_tolerance=1, y_tolerance=3) or ""
    except Exception:
        return ""


_PAGE_WINDOWS = {
    2015: [(538, 543)],
    2016: [(541, 568)],
    2017: [(605, 630)],
    2018: [(990, 1000)],
    2019: [(1, 20)],
    2020: [(266, 275)],
    2023: [(231, 259)],
    2021: [(77, 90)],
    2022: [(218, 226)],
    2024: [(229, 237)],
    2025: [(1, 20)],
    2026: [(229, 237)],
}

_ORG_PROGRAMA_WINDOWS = {
    # FIX 2026-07-27: ventana ampliada de (230,523) a (230,628). El límite
    # anterior cortaba la Sección 20 "POLÍTICAS SOCIALES, FAMILIA, IGUALDAD Y
    # JUSTICIA" (empieza p.567), que contiene dependencia/discapacidad/
    # diversidad/igualdad — esos 4 conceptos caían a 0 o casi 0 en 2018.
    # p.628 es la ÚLTIMA página del informe "Detalle Orgánico/Económico de
    # los Programas"; p.629 es la portadilla del informe SIGUIENTE ("...por
    # Nivel de Especificación", empieza p.631) que repite las MISMAS
    # secciones en otro agrupamiento — incluirlo duplicaría el gasto, por
    # eso el corte en 628 es intencional, no arbitrario. Verificado: 2018
    # pasa de 53 a 72 filas, 9 a 12 conceptos, dependencia 9,68→59,68 M€
    # (encaja con 2017=58,46 y 2019=58,82), discapacidad 0→22,10 M€ (2017=
    # 22,29, 2019=22,66), total 1.303→1.485 M€ (entre 2017=1.453 y 2019=
    # 1.533, sin inflación). Ver logs/progreso.md 2026-07-27.
    2018: [(230, 628)],
}


def _allowed_page(anio: int, p_idx: int, total_pages: int) -> bool:
    windows = _PAGE_WINDOWS.get(anio)
    if not windows:
        return True
    if total_pages < min(start for start, _ in windows):
        return True
    return any(start <= p_idx <= min(end, total_pages) for start, end in windows)


def _parse_resumen_programas(pdfs: list[Path], anio: int) -> list[dict]:
    """Parsea tablas `Resumen ... por Programas` o `Resumen de Sección`.

    Es la vía más estable para los tomos 2015-2018: cada fila trae `NNNN -
    denominación` y varias columnas económicas, con el total al final.
    """
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        return []

    rows: list[dict] = []
    seen: set[str] = set()
    for pdf_path in pdfs:
        try:
            pdf = pdfplumber.open(pdf_path)
        except Exception:
            continue
        with pdf:
            for p_idx, page in enumerate(pdf.pages, start=1):
                if not _allowed_page(anio, p_idx, len(pdf.pages)):
                    continue
                txt = _page_text(page)
                if "Anexo de Inversiones" in txt or "Proyectos de gastos" in txt:
                    continue
                if "Resumen" not in txt and "Programa" not in txt:
                    continue
                pending: dict | None = None
                for line in txt.splitlines():
                    m = _RE_CODE4.match(line)
                    if m:
                        pending = {
                            "code": m.group("code"),
                            "parts": [m.group("rest")],
                        }
                    elif pending is not None:
                        if line.strip().lower().startswith("total "):
                            pending = None
                            continue
                        pending["parts"].append(line)
                    if pending is None:
                        continue
                    text = " ".join(pending["parts"])
                    vals = _amounts(text)
                    total = _row_total(vals)
                    if total is None:
                        continue
                    first_amount = _RE_AMOUNT.search(text)
                    denom = _norm_label(text[:first_amount.start()] if first_amount else text)
                    code = _code4_to_dotted(pending["code"])
                    if code not in seen:
                        seen.add(code)
                        rows.append(make_row(pagina=p_idx, codigo=code,
                                             denominacion=denom,
                                             importe=float(total), anio=anio))
                    pending = None
    return rows


def _parse_resumen_programas_invertido(pdfs: list[Path], anio: int) -> list[dict]:
    """Parsea resúmenes 2015 cuyo texto sale como tokens invertidos.

    El bloque `Resumen Gastos por Capítulos y Programas` del BOLR 2015 expone
    cada celda como una línea independiente y con caracteres en orden inverso:
    diez importes, código `NNNN`, guion y denominación. Revertimos los tokens y
    reconstruimos una fila por programa usando el último importe como total.
    """
    if anio != 2015:
        return []
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        return []

    def is_amount_token(tok: str) -> bool:
        # Los códigos de nivel 1-9 aparecen como tokens sueltos; no son importes.
        if tok in {"0", "o", "O"}:
            return True
        if re.fullmatch(r"-?\d", tok):
            return False
        return _to_int_eur(tok) is not None

    header_noise = {
        "ADMINISTRACIÑN", "ADMINISTRACIÓN", "GENERAL", "Total", "Programa",
        "PASIVOS", "ACTIVOS", "TRANSFERENCIA", "TRANSFERENCIAS", "INVERSIONES",
        "FONDO", "GASTOS", "CORRIENTES", "SERVICIOS", "FINANCIEROS", "PERSONAL",
        "AMORTIZACION", "AMORTIZACIÓN", "BIENES", "CAPITAL", "REALES",
        "CONTINGENCIA",
    }

    def keep_word(tok: str) -> bool:
        if not tok or tok == "-":
            return False
        if tok.upper() in {"GOBIERNO", "DE", "LA", "RIOJA"}:
            return False
        if len(tok) == 1:
            return False
        if tok in header_noise:
            return False
        if "Consejer" in tok or "Hacienda" in tok:
            return False
        if "Resumen" in tok or "Programas" in tok or "Cap" in tok:
            return False
        return True

    rows: list[dict] = []
    seen: set[str] = set()
    for pdf_path in pdfs:
        try:
            pdf = pdfplumber.open(pdf_path)
        except Exception:
            continue
        with pdf:
            for p_idx, page in enumerate(pdf.pages, start=1):
                if not _allowed_page(anio, p_idx, len(pdf.pages)):
                    continue
                txt = _page_text(page)
                if "nemuseR" not in txt and "Resumen" not in txt:
                    continue
                tokens = [_reverse_pdf_token(line)
                          for line in reversed(txt.splitlines()) if line.strip()]
                for i in range(0, max(0, len(tokens) - 12)):
                    if not all(is_amount_token(tokens[i + j]) for j in range(10)):
                        continue
                    if not re.fullmatch(r"\d{4}", tokens[i + 10] or ""):
                        continue
                    if tokens[i + 11] != "-":
                        continue
                    total = _to_int_eur(tokens[i + 9])
                    if total is None or total <= 0:
                        continue
                    code = _code4_to_dotted(tokens[i + 10])
                    if code in seen:
                        continue
                    words: list[str] = []
                    j = i + 12
                    while j < len(tokens):
                        if j + 2 < len(tokens) and all(is_amount_token(tokens[j + k]) for k in range(3)):
                            break
                        if re.fullmatch(r"\d{4}", tokens[j] or ""):
                            break
                        if keep_word(tokens[j]):
                            words.append(tokens[j])
                        j += 1
                    denom = _norm_label(" ".join(words))
                    seen.add(code)
                    rows.append(make_row(pagina=p_idx, codigo=code,
                                         denominacion=denom,
                                         importe=float(total), anio=anio))
    return rows


def _parse_detalle_organico_programas(pdfs: list[Path], anio: int) -> list[dict]:
    """Parsea encabezados `Programa NNNN denominación total`.

    Es la vía estable para 2018, cuyo resumen funcional-económico aparece
    invertido en el PDF pero cuyo detalle orgánico/económico conserva el
    encabezado de cada programa como texto legible.
    """
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        return []

    windows = _ORG_PROGRAMA_WINDOWS.get(anio)
    if not windows:
        return []

    agg: dict[str, float] = {}
    denom_by: dict[str, str] = {}
    page_by: dict[str, int] = {}
    for pdf_path in pdfs:
        try:
            pdf = pdfplumber.open(pdf_path)
        except Exception:
            continue
        with pdf:
            for p_idx, page in enumerate(pdf.pages, start=1):
                if not any(start <= p_idx <= min(end, len(pdf.pages)) for start, end in windows):
                    continue
                txt = _page_text(page)
                if "Detalle Orgánico / Económico de los Programas" not in txt:
                    continue
                for line in txt.splitlines():
                    m = _RE_ORG_PROGRAMA.match(line)
                    if not m:
                        continue
                    code = _code4_to_dotted(m.group("code"))
                    importe = parse_eur(m.group("importe"))
                    if importe != importe:
                        continue
                    agg[code] = agg.get(code, 0.0) + importe
                    denom_by.setdefault(code, _norm_label(m.group("denom")))
                    page_by.setdefault(code, p_idx)

    return [
        make_row(pagina=page_by[c], codigo=c,
                 denominacion=denom_by[c], importe=agg[c], anio=anio)
        for c in sorted(agg.keys())
    ]


def _parse_informe_resumen_funcional(pdfs: list[Path], anio: int) -> list[dict]:
    """Parsea `Informe Resumen General Funcional - Económico`.

    Este formato aparece, al menos, en 2020: las columnas no traen guiones
    (`Funcio Subfuncion Programa`) y el grupo presupuestario se deduce por
    bloques hasta cada línea `Total Grupo N`.
    """
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        return []

    next_group = {1: 2, 2: 3, 3: 4, 4: 9, 9: None}
    rows: list[dict] = []
    seen: set[str] = set()
    for pdf_path in pdfs:
        try:
            pdf = pdfplumber.open(pdf_path)
        except Exception:
            continue
        grupo: int | None = 1
        funcion = subfun = prog = None
        with pdf:
            for p_idx, page in enumerate(pdf.pages, start=1):
                if not _allowed_page(anio, p_idx, len(pdf.pages)):
                    continue
                txt = _page_text(page)
                if "Informe Resumen General Funcional" not in txt:
                    continue
                for line in txt.splitlines():
                    compact = re.sub(r"\s+", " ", line).strip()
                    m_total_grupo = re.search(r"Total Grupo\s+(\d)", compact)
                    if m_total_grupo:
                        grupo = next_group.get(int(m_total_grupo.group(1)))
                        continue
                    if compact.startswith("Total "):
                        continue
                    matches = _amount_matches(line)
                    if not matches:
                        continue
                    amount_start = matches[0][0].start()
                    head = line[:amount_start]
                    segments = _informe_segments(head)
                    if not segments or grupo is None:
                        continue
                    if len(segments) >= 3:
                        funcion = int(segments[-3][0])
                        subfun = int(segments[-2][0])
                        prog = int(segments[-1][0])
                        denom = segments[-1][1]
                    elif len(segments) == 2 and funcion is not None:
                        subfun = int(segments[-2][0])
                        prog = int(segments[-1][0])
                        denom = segments[-1][1]
                    elif len(segments) == 1 and funcion is not None and subfun is not None:
                        prog = int(segments[-1][0])
                        denom = segments[-1][1]
                    else:
                        continue
                    vals = [v for _, v in matches]
                    total = vals[-1]
                    if total <= 0:
                        continue
                    code = f"{grupo}.{funcion}.{subfun}.{prog}"
                    if code in seen:
                        continue
                    seen.add(code)
                    rows.append(make_row(pagina=p_idx, codigo=code,
                                         denominacion=denom,
                                         importe=float(total), anio=anio))
    return rows


def _parse_detalle_funcional(pdfs: list[Path], anio: int) -> list[dict]:
    """Parsea `Detalle Gastos Funcional / Económico` desde texto pdfplumber.

    La tabla aparece con cuatro niveles jerárquicos (grupo, función, subfunción,
    programa). El texto extraído no conserva columnas de forma perfecta, así que
    se usa una máquina de estado por niveles y los importes de capítulo/total.
    """
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        return []

    rows: list[dict] = []
    seen: set[str] = set()
    for pdf_path in pdfs:
        try:
            pdf = pdfplumber.open(pdf_path)
        except Exception:
            continue
        grupo = funcion = subfun = prog = None
        with pdf:
            for p_idx, page in enumerate(pdf.pages, start=1):
                if not _allowed_page(anio, p_idx, len(pdf.pages)):
                    continue
                txt = _page_text(page)
                is_detail_page = (
                    "Detalle Gastos Funcional" in txt
                    or "Funcional / Econ" in txt
                    or "Funcional I Econ" in txt
                    or "Grupo Función Subfunción" in txt
                    or anio in _PAGE_WINDOWS
                    or (len(pdf.pages) <= 20 and "Total Programa" in txt)
                )
                if not is_detail_page:
                    continue
                has_total_col = bool(re.search(r"Cap\.\s*9\s+Total", txt))
                for line in txt.splitlines():
                    compact = re.sub(r"\s+", "", line).lower()
                    if "totalfunc" in compact or "totalsubfunc" in compact or "totalgrupo" in compact:
                        continue
                    segment_matches = list(_RE_SEGMENT.finditer(line))
                    segments = [m.group(1) for m in segment_matches]
                    if len(segments) >= 4:
                        grupo, funcion, subfun, prog = segments[-4:]
                    elif len(segments) == 3 and grupo is not None:
                        funcion, subfun, prog = segments
                    elif len(segments) == 2 and grupo is not None and funcion is not None:
                        subfun, prog = segments
                    elif len(segments) == 1 and all(v is not None for v in (grupo, funcion, subfun)):
                        prog = segments[0]
                    amount_start = segment_matches[-1].end() if segment_matches else 0
                    if "totalprograma" in compact:
                        m_total = re.search(r"total\s*programa", line, flags=re.I)
                        if m_total:
                            amount_start = max(amount_start, m_total.end())
                    vals = _amounts(line[amount_start:])
                    if not vals:
                        continue
                    head = line[:amount_start]
                    if not all(v is not None for v in (grupo, funcion, subfun, prog)):
                        continue
                    if "totalprograma" in compact or "totalprogram" in compact or segments:
                        total = _row_total(vals) if has_total_col else (sum(vals) if vals else None)
                        if total is None or total <= 0:
                            continue
                        code = f"{grupo}.{funcion}.{subfun}.{prog}"
                        if code in seen:
                            continue
                        seen.add(code)
                        denom = _norm_label(re.sub(_RE_SEGMENT, " ", head))
                        denom = re.sub(r"\bTotal\s*Programa\b", "", denom, flags=re.I).strip()
                        rows.append(make_row(pagina=p_idx, codigo=code,
                                             denominacion=denom,
                                             importe=float(total), anio=anio))
    return rows


def _supplement_image_only_rows(rows: list[dict], anio: int) -> list[dict]:
    """Añade filas visibles en páginas que pdfminer no expone como texto.

    En la Ley BOLR 2024, las páginas 23300 y 23304 renderizan tablas limpias
    de `Detalle Gastos Funcional / Económico`, pero internamente solo exponen
    la cabecera del boletín a pdfplumber/pdfminer. Los importes siguientes se
    han transcrito de esas dos páginas renderizadas y cierran el Total General
    de Gastos de la misma tabla.
    """
    if anio != 2024 or not rows:
        return rows

    existing = {r.get("codigo") for r in rows}
    manual = [
        # Página BOLR 23300: grupo 3, educación/cultura/deporte.
        (232, "3.2.1.1", "ADMINISTRACION GENERAL DE EDUCACION", 17_282_751),
        (232, "3.2.2.1", "ENSENANZA REGIMEN GENERAL", 298_487_862),
        (232, "3.2.2.2", "ENSENANZA REGIMEN ESPECIAL", 13_677_332),
        (232, "3.2.2.3", "FORMACION PROFESIONAL", 43_249_892),
        (232, "3.2.2.5", "ENSENANZA UNIVERSITARIA", 30_461_002),
        (232, "3.3.1.0", "ADMINISTRACION GENERAL DE CULTURA", 749_800),
        (232, "3.3.1.1", "PROMOCION DE LA CULTURA", 8_422_372),
        (232, "3.3.1.2", "MUSEOS ARCHIVOS Y BIBLIOTECAS", 4_951_825),
        (232, "3.3.2.1", "PATRIMONIO HISTORICO-ARTISTICO", 4_952_718),
        (232, "3.4.1.1", "DEPORTE", 10_005_824),
        # Página BOLR 23304: cierre del grupo 9, administración/deuda.
        (236, "9.2.3.1", "ADMINISTRACION Y GESTION TRIBUTARIA", 12_264_577),
        (236, "9.2.4.1", "CONTROL INTERNO AUDITORIA Y CONTABILIDAD", 3_271_755),
        (236, "9.2.5.1", "ESTADISTICA", 1_156_078),
        (236, "9.9.1.1", "FONDO DE CONTINGENCIA DE EJECUCION PRESUPUESTARIA", 1_406_820),
        (236, "9.3.1.1", "SERVICIOS A ENTIDADES LOCALES", 3_607_598),
        (236, "9.5.1.1", "AMORTIZACION Y GASTOS FINANCIEROS", 176_892_127),
    ]
    for pagina, codigo, denominacion, importe in manual:
        if codigo in existing:
            continue
        rows.append(make_row(pagina=pagina, codigo=codigo,
                             denominacion=denominacion,
                             importe=float(importe), anio=anio))
        existing.add(codigo)
    return rows


def extract(input_path: Path, anio: int) -> ExtractionResult:
    # 1) XLSX adjunto manda
    rows = _from_xlsx(input_path, anio)
    if rows:
        return ExtractionResult(rows=rows, motor="lar-xlsx-adjunto")

    pdfs = _iter_input_pdfs(input_path, anio)

    # 2) Detalle orgánico/económico por programa (2018)
    rows = _parse_detalle_organico_programas(pdfs, anio)
    if rows:
        return ExtractionResult(rows=rows, motor="lar-detalle-organico-programas")

    # 3) Resumen por programas invertido (BOLR 2015)
    rows = _parse_resumen_programas_invertido(pdfs, anio)
    if rows:
        return ExtractionResult(rows=rows, motor="lar-resumen-programas-invertido")

    # 4) Resumen por programas (código 4 dígitos + total)
    rows = _parse_resumen_programas(pdfs, anio)
    if rows:
        return ExtractionResult(rows=rows, motor="lar-resumen-programas-pdfplumber")

    # 5) Informe resumen funcional-económico (formato sin guiones, p.ej. 2020)
    rows = _parse_informe_resumen_funcional(pdfs, anio)
    if rows:
        return ExtractionResult(rows=rows, motor="lar-informe-resumen-funcional")

    # 6) PDF "Detalle Funcional/Económico" vía texto pdfplumber
    rows = _parse_detalle_funcional(pdfs, anio)
    rows = _supplement_image_only_rows(rows, anio)
    motor = "lar-pdfplumber-funcional-economico" if rows else "lar-pendiente"
    notes = "" if rows else (
        "No se localizaron filas de resumen/programa ni detalle funcional."
    )
    return ExtractionResult(rows=rows, motor=motor, notes=notes)
